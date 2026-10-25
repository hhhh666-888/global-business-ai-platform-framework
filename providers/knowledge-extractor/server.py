import base64
import json
import os
import re
import shutil
import signal
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile
from email import policy
from email.parser import BytesParser
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

DATA_ROOT = Path(os.environ.get("KNOWLEDGE_EXTRACTOR_DATA", "/data"))
PORT = int(os.environ.get("KNOWLEDGE_EXTRACTOR_PORT", "8093"))
MAX_BYTES = int(os.environ.get("KNOWLEDGE_MAX_DOCUMENT_BYTES", str(25 * 1024 * 1024)))
OCR_LANGUAGES = os.environ.get("KNOWLEDGE_OCR_LANGUAGES", "eng+chi_sim")
MAX_NORMALIZED_CHARS = int(os.environ.get("KNOWLEDGE_MAX_NORMALIZED_CHARS", "8000000"))
MAX_ARCHIVE_MEMBERS = int(os.environ.get("KNOWLEDGE_MAX_ARCHIVE_MEMBERS", "20000"))
MAX_ARCHIVE_UNCOMPRESSED_BYTES = int(os.environ.get("KNOWLEDGE_MAX_ARCHIVE_UNCOMPRESSED_BYTES", str(512 * 1024 * 1024)))
MAX_ARCHIVE_MEMBER_BYTES = int(os.environ.get("KNOWLEDGE_MAX_ARCHIVE_MEMBER_BYTES", str(128 * 1024 * 1024)))
MAX_ARCHIVE_COMPRESSION_RATIO = int(os.environ.get("KNOWLEDGE_MAX_ARCHIVE_COMPRESSION_RATIO", "500"))
OPENDATALOADER_TIMEOUT_SEC = int(os.environ.get("KNOWLEDGE_OPENDATALOADER_TIMEOUT_SEC", "600"))
OCR_TIMEOUT_SEC = int(os.environ.get("KNOWLEDGE_OCR_TIMEOUT_SEC", "1200"))
OFFICE_TIMEOUT_SEC = int(os.environ.get("KNOWLEDGE_OFFICE_TIMEOUT_SEC", "240"))
IMAGE_OCR_TIMEOUT_SEC = int(os.environ.get("KNOWLEDGE_IMAGE_OCR_TIMEOUT_SEC", "300"))
HANDED_OFF_RETENTION_DAYS = int(os.environ.get("KNOWLEDGE_HANDED_OFF_RETENTION_DAYS", "7"))
FAILED_RETENTION_DAYS = int(os.environ.get("KNOWLEDGE_FAILED_RETENTION_DAYS", "30"))
MAINTENANCE_INTERVAL_SEC = int(os.environ.get("KNOWLEDGE_MAINTENANCE_INTERVAL_SEC", "3600"))
PARSE_TIMEOUT_SEC = int(os.environ.get("KNOWLEDGE_PARSE_TIMEOUT_SEC", "2400"))
MAX_PARSER_RSS_BYTES = int(os.environ.get("KNOWLEDGE_MAX_PARSER_RSS_BYTES", str(2 * 1024 * 1024 * 1024)))
MAX_PARSER_FILE_BYTES = int(os.environ.get("KNOWLEDGE_MAX_PARSER_FILE_BYTES", str(256 * 1024 * 1024)))
MAX_PARSER_TEMP_BYTES = int(os.environ.get("KNOWLEDGE_MAX_PARSER_TEMP_BYTES", str(1024 * 1024 * 1024)))
MAX_SPREADSHEET_ROWS = int(os.environ.get("KNOWLEDGE_MAX_SPREADSHEET_ROWS", "100000"))
MAX_SPREADSHEET_CELLS = int(os.environ.get("KNOWLEDGE_MAX_SPREADSHEET_CELLS", "2000000"))
if min(PARSE_TIMEOUT_SEC, MAX_PARSER_RSS_BYTES, MAX_PARSER_FILE_BYTES, MAX_PARSER_TEMP_BYTES, MAX_SPREADSHEET_CELLS, MAX_SPREADSHEET_ROWS) <= 0:
    raise ValueError("invalid_parser_resource_limits")
PARSER_CHILD = False
JOBS_ROOT = DATA_ROOT / "jobs"
DB_PATH = DATA_ROOT / "extractor.db"


def ensure_storage():
    DATA_ROOT.mkdir(parents=True, exist_ok=True)
    JOBS_ROOT.mkdir(parents=True, exist_ok=True)
WAKE = threading.Event()
STOP = threading.Event()
ACTIVE_JOBS = {}
ACTIVE_JOBS_LOCK = threading.Lock()
CURRENT_JOB = threading.local()

SUPPORTED = {
    ".pdf", ".docx", ".xlsx", ".pptx", ".docm", ".xlsm", ".pptm", ".txt", ".md", ".csv", ".json", ".xml",
    ".html", ".htm", ".doc", ".xls", ".ppt", ".rtf", ".odt", ".ods", ".odp",
    ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp",
    ".yaml", ".yml", ".log", ".tsv", ".eml", ".epub",
}
LEGACY_TARGET = {
    ".doc": "docx", ".docm": "docx", ".rtf": "docx", ".odt": "docx",
    ".xls": "xlsx", ".xlsm": "xlsx", ".ods": "xlsx",
    ".ppt": "pptx", ".pptm": "pptx", ".odp": "pptx",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm", ".yaml", ".yml", ".log", ".tsv"}
ARCHIVE_EXTENSIONS = {".docx", ".xlsx", ".pptx", ".docm", ".xlsm", ".pptm", ".odt", ".ods", ".odp", ".epub"}


def now():
    return datetime.now(timezone.utc).isoformat()


def db():
    ensure_storage()
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA busy_timeout=30000")
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS jobs (
          id TEXT PRIMARY KEY,
          state TEXT NOT NULL,
          stage TEXT NOT NULL,
          source_filename TEXT NOT NULL,
          media_type TEXT NOT NULL,
          input_path TEXT NOT NULL,
          result_path TEXT,
          parser TEXT,
          metrics_json TEXT,
          error TEXT,
          handoff_external_id TEXT,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        )
        """
    )
    conn.commit()
    return conn


def reset_interrupted_jobs():
    conn = db()
    conn.execute("UPDATE jobs SET state='QUEUED',stage='RECOVERED',updated_at=? WHERE state='RUNNING'", (now(),))
    conn.commit()
    conn.close()


def purge_expired_jobs():
    now_dt = datetime.now(timezone.utc)
    handed_off_before = (now_dt - timedelta(days=HANDED_OFF_RETENTION_DAYS)).isoformat()
    failed_before = (now_dt - timedelta(days=FAILED_RETENTION_DAYS)).isoformat()
    conn = db()
    rows = conn.execute(
        """
        SELECT id FROM jobs
        WHERE (state='READY' AND handoff_external_id IS NOT NULL AND updated_at < ?)
           OR (state IN ('FAILED','CANCELLED') AND updated_at < ?)
        """,
        (handed_off_before, failed_before),
    ).fetchall()
    if rows:
        conn.executemany("DELETE FROM jobs WHERE id=?", [(row["id"],) for row in rows])
        conn.commit()
    conn.close()
    for row in rows:
        shutil.rmtree(JOBS_ROOT / row["id"], ignore_errors=True)
    return len(rows)


def run_bounded(command, *, timeout, label, accepted_codes=(0,)):
    cancel_event = getattr(CURRENT_JOB, "cancel_event", None)
    if cancel_event and cancel_event.is_set():
        raise RuntimeError("extraction_cancelled")
    # Disk-backed logs avoid communicate() retaining unlimited converter output.
    # Parser children inherit the supervisor group, including converter children.
    # Visible private files let the parent account for converter log disk usage.
    with tempfile.NamedTemporaryFile() as output, tempfile.NamedTemporaryFile() as errors:
        proc = subprocess.Popen(command, stdout=output, stderr=errors, start_new_session=not PARSER_CHILD)
        deadline = time.monotonic() + timeout
        try:
            while True:
                cancelled = cancel_event is not None and cancel_event.is_set()
                remaining = deadline - time.monotonic()
                if cancelled or remaining <= 0 or STOP.is_set():
                    if PARSER_CHILD:
                        proc.kill()
                    else:
                        try:
                            os.killpg(proc.pid, signal.SIGKILL)
                        except ProcessLookupError:
                            pass
                    proc.wait()
                    if cancelled or STOP.is_set():
                        raise RuntimeError("extraction_cancelled")
                    raise RuntimeError(f"{label}_timeout:{timeout}s")
                try:
                    proc.wait(timeout=min(remaining, 0.5))
                    break
                except subprocess.TimeoutExpired:
                    continue
            errors.seek(max(0, errors.tell() - 8000))
            stderr = errors.read(8000).decode("utf-8", errors="replace")
            if proc.returncode == -signal.SIGXFSZ:
                raise RuntimeError("document_parse_storage_limit")
            if proc.returncode not in accepted_codes:
                raise RuntimeError(f"{label}_failed:exit={proc.returncode}:{clean_text(stderr)[-2000:]}")
            if output.tell() > MAX_NORMALIZED_CHARS * 4:
                raise RuntimeError(f"{label}_output_too_large")
            output.seek(0)
            return output.read().decode("utf-8", errors="replace"), stderr
        finally:
            if proc.poll() is None:
                proc.kill()
                proc.wait()


def parser_group_rss(group_id):
    total = 0
    for entry in Path("/proc").iterdir():
        if not entry.name.isdigit():
            continue
        try:
            fields = (entry / "stat").read_text().rsplit(")", 1)[1].split()
            if int(fields[2]) == group_id:
                total += int((entry / "statm").read_text().split()[1]) * os.sysconf("SC_PAGE_SIZE")
        except (OSError, ValueError, IndexError):
            continue
    return total


def run_isolated_parser(command, workdir, *, timeout=PARSE_TIMEOUT_SEC, max_rss=MAX_PARSER_RSS_BYTES, max_temp=MAX_PARSER_TEMP_BYTES):
    cancel_event = getattr(CURRENT_JOB, "cancel_event", None)
    if STOP.is_set() or (cancel_event and cancel_event.is_set()):
        raise RuntimeError("extraction_cancelled")
    env = {**os.environ, "TMPDIR": str(Path(workdir) / "tmp")}
    Path(env["TMPDIR"]).mkdir()
    with (Path(workdir) / "stdout.log").open("wb") as output, (Path(workdir) / "stderr.log").open("wb") as errors:
        proc = subprocess.Popen(command, stdout=output, stderr=errors, env=env, start_new_session=True)
        deadline = time.monotonic() + timeout
        next_resource_check = 0
        try:
            while True:
                if STOP.is_set() or (cancel_event and cancel_event.is_set()):
                    raise RuntimeError("extraction_cancelled")
                if time.monotonic() >= deadline:
                    raise RuntimeError(f"document_parse_timeout:{timeout}s")
                # WNOWAIT preserves the leader PID until group cleanup, preventing
                # a recycled process/group ID from being signalled accidentally.
                if os.waitid(os.P_PID, proc.pid, os.WEXITED | os.WNOHANG | os.WNOWAIT):
                    break
                if time.monotonic() >= next_resource_check:
                    next_resource_check = time.monotonic() + 0.5
                    if parser_group_rss(proc.pid) > max_rss:
                        raise RuntimeError("document_parse_memory_limit")
                    size = 0
                    for index, path in enumerate(Path(workdir).rglob("*"), 1):
                        if index > 50000:
                            raise RuntimeError("document_parse_storage_limit")
                        try:
                            if path.is_file():
                                size += path.stat().st_size
                        except FileNotFoundError:
                            continue  # converters remove their scratch files concurrently
                        if size > max_temp:
                            raise RuntimeError("document_parse_storage_limit")
                time.sleep(0.05)
        finally:
            try:
                os.killpg(proc.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
            proc.wait()
    return proc.returncode


def parse_document_isolated(path):
    with tempfile.TemporaryDirectory(prefix="gbap-parser-") as root:
        result = Path(root) / "result.json"
        status = run_isolated_parser([sys.executable, str(Path(__file__).resolve()), "--parse-document", str(path.resolve()), str(result)], root)
        if status == -signal.SIGXFSZ:
            raise RuntimeError("document_parse_storage_limit")
        if not result.exists() or result.stat().st_size > MAX_NORMALIZED_CHARS * 4 + 1024 * 1024:
            raise RuntimeError("document_parser_result_invalid")
        value = json.loads(result.read_text("utf-8"))
        if status != 0 or value.get("error"):
            raise RuntimeError(value.get("error") or "document_parser_failed")
        return value["text"], value["parser"], value["metrics"], value["attempts"]


def parser_child_main(input_path, result_path):
    global PARSER_CHILD
    PARSER_CHILD = True
    import resource
    resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_PARSER_FILE_BYTES, MAX_PARSER_FILE_BYTES))
    try:
        text, parser, metrics, attempts = parse_document(Path(input_path))
        if len(text) > MAX_NORMALIZED_CHARS:
            raise RuntimeError("normalized_content_too_large")
        value = {"text": text, "parser": parser, "metrics": metrics, "attempts": attempts}
    except Exception as exc:
        Path(result_path).write_text(json.dumps({"error": str(exc)[:8000]}, ensure_ascii=False), "utf-8")
        return 2
    Path(result_path).write_text(json.dumps(value, ensure_ascii=False), "utf-8")
    return 0


def validate_archive_container(path):
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise RuntimeError(f"archive_too_many_members:{len(members)}>{MAX_ARCHIVE_MEMBERS}")
            total = 0
            for item in members:
                name = item.filename.replace("\\", "/")
                parts = [part for part in name.split("/") if part not in ("", ".")]
                if name.startswith("/") or any(part == ".." for part in parts):
                    raise RuntimeError("archive_unsafe_path")
                if item.flag_bits & 0x1:
                    raise RuntimeError("archive_encrypted_member")
                if item.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise RuntimeError(f"archive_member_too_large:{item.file_size}>{MAX_ARCHIVE_MEMBER_BYTES}")
                total += item.file_size
                if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise RuntimeError(f"archive_uncompressed_too_large:{total}>{MAX_ARCHIVE_UNCOMPRESSED_BYTES}")
                if item.file_size and item.compress_size > 0:
                    ratio = item.file_size / item.compress_size
                    if ratio > MAX_ARCHIVE_COMPRESSION_RATIO:
                        raise RuntimeError(f"archive_compression_ratio_too_high:{ratio:.1f}>{MAX_ARCHIVE_COMPRESSION_RATIO}")
    except zipfile.BadZipFile as exc:
        raise RuntimeError("invalid_archive_container") from exc


def safe_name(value):
    value = Path(value or "document.bin").name
    return re.sub(r"[^A-Za-z0-9._()\-\u4e00-\u9fff]+", "_", value)[:180] or "document.bin"


def clean_text(text):
    text = (text or "").replace("\x00", "")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def tokenish_count(text):
    return len(re.findall(r"[A-Za-z0-9][A-Za-z0-9_./:+\-]*|[\u3400-\u9fff]", text or ""))


def quality(text, *, page_count=None, pages_with_text=None):
    text = clean_text(text)
    non_ws = len(re.sub(r"\s+", "", text))
    tokens = tokenish_count(text)
    replacement_rate = text.count("\ufffd") / max(1, len(text))
    coverage = None
    if page_count and page_count > 0 and pages_with_text is not None:
        coverage = pages_with_text / page_count
    reasons = []
    if non_ws < 40:
        reasons.append("too_few_characters")
    if tokens < 8:
        reasons.append("too_few_tokens")
    if replacement_rate > 0.02:
        reasons.append("encoding_noise")
    if page_count and page_count >= 4 and coverage is not None and coverage < 0.70:
        reasons.append("low_page_coverage")
    return {
        "accepted": not reasons,
        "characters": len(text),
        "nonWhitespaceCharacters": non_ws,
        "tokenishCount": tokens,
        "replacementRate": round(replacement_rate, 6),
        "pageCount": page_count,
        "pagesWithText": pages_with_text,
        "pageCoverage": round(coverage, 4) if coverage is not None else None,
        "reasons": reasons,
    }


def decode_text_file(path):
    from charset_normalizer import from_bytes
    raw = Path(path).read_bytes()
    best = from_bytes(raw).best()
    return str(best) if best else raw.decode("utf-8", errors="replace")


def markdown_table(rows):
    rows = [[("" if value is None else str(value)).replace("|", "\\|").replace("\n", "<br>") for value in row] for row in rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    body = rows[1:]
    out = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    out.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(out)


def extract_docx(path):
    from docx import Document
    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        text = clean_text(p.text)
        if not text:
            continue
        style = (p.style.name if p.style else "").lower()
        match = re.search(r"heading\s*(\d+)", style)
        if match:
            out.append("#" * max(1, min(6, int(match.group(1)))) + " " + text)
        else:
            out.append(text)
    for index, table in enumerate(doc.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            out.append(f"\n### Table {index}\n\n{markdown_table(rows)}")
    return clean_text("\n\n".join(out)), {"tables": len(doc.tables)}


def extract_xlsx(path):
    from openpyxl import load_workbook
    formulas = load_workbook(path, read_only=True, data_only=False)
    values = None
    try:
        values = load_workbook(path, read_only=True, data_only=True)
        out = []
        sheet_metrics = []
        value_sheets = {ws.title: ws for ws in values.worksheets}
        cells_seen = rows_seen = characters = 0
        for ws in formulas.worksheets:
            value_ws = value_sheets.get(ws.title)
            # A producer can claim the entire Excel grid while storing two rows.
            # Follow actual XML rows rather than padding to the declared dimension.
            ws.reset_dimensions()
            if value_ws:
                value_ws.reset_dimensions()
            rows = []
            value_rows = value_ws.iter_rows(values_only=True) if value_ws else None
            for formula_row in ws.iter_rows():
                cells_seen += len(formula_row)
                rows_seen += 1
                if cells_seen > MAX_SPREADSHEET_CELLS or rows_seen > MAX_SPREADSHEET_ROWS:
                    raise RuntimeError("document_complexity_limit")
                cached_row = next(value_rows, ()) if value_rows is not None else ()
                rendered = []
                for index, cell in enumerate(formula_row):
                    raw = cell.value
                    cached = cached_row[index] if index < len(cached_row) else None
                    if isinstance(raw, str) and raw.startswith("=") and cached not in (None, ""):
                        rendered.append(f"{raw} => {cached}")
                    else:
                        rendered.append("" if raw is None else str(raw))
                while rendered and not rendered[-1].strip():
                    rendered.pop()
                if any(value.strip() for value in rendered):
                    characters += sum(len(value) + 4 for value in rendered)
                    if characters > MAX_NORMALIZED_CHARS:
                        raise RuntimeError("normalized_content_too_large")
                    rows.append(rendered)
            if not rows:
                continue
            min_width = min((next((i for i, v in enumerate(row) if v.strip()), len(row)) for row in rows), default=0)
            if min_width:
                rows = [row[min_width:] for row in rows]
            out.append(f"## Sheet: {ws.title}\n\n{markdown_table(rows)}")
            sheet_metrics.append({"name": ws.title, "rows": len(rows), "columns": max((len(r) for r in rows), default=0)})
        return clean_text("\n\n".join(out)), {"sheets": sheet_metrics}
    finally:
        formulas.close()
        if values is not None:
            values.close()


def extract_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for number, slide in enumerate(prs.slides, 1):
        out.append(f"## Slide {number}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = clean_text(shape.text)
                if text:
                    out.append(text)
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                if rows:
                    out.append(markdown_table(rows))
        try:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes:
                out.append(f"### Notes\n\n{notes}")
        except Exception:
            pass
    return clean_text("\n\n".join(out)), {"slides": len(prs.slides)}


def extract_html(path):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(decode_text_file(path), "html.parser")
    for node in soup(["script", "style", "noscript"]):
        node.decompose()
    lines = []
    title = soup.title.string.strip() if soup.title and soup.title.string else ""
    if title:
        lines.append("# " + title)
    for node in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "th", "td"]):
        text = clean_text(node.get_text(" ", strip=True))
        if not text:
            continue
        if node.name.startswith("h"):
            lines.append("#" * int(node.name[1]) + " " + text)
        elif node.name == "li":
            lines.append("- " + text)
        else:
            lines.append(text)
    return clean_text("\n\n".join(lines)), {}


def extract_eml(path):
    message = BytesParser(policy=policy.default).parsebytes(Path(path).read_bytes())
    out = []
    for label, value in [
        ("Subject", message.get("subject")),
        ("From", message.get("from")),
        ("To", message.get("to")),
        ("Date", message.get("date")),
    ]:
        if value:
            out.append(f"{label}: {value}")
    body = message.get_body(preferencelist=("plain", "html"))
    if body:
        content = body.get_content()
        if body.get_content_type() == "text/html":
            from bs4 import BeautifulSoup
            content = BeautifulSoup(content, "html.parser").get_text("\n", strip=True)
        out.append(clean_text(content))
    return clean_text("\n\n".join(out)), {}


def extract_epub(path):
    from bs4 import BeautifulSoup
    out = []
    with zipfile.ZipFile(path) as archive:
        names = sorted(name for name in archive.namelist() if name.lower().endswith((".xhtml", ".html", ".htm")))
        for name in names:
            soup = BeautifulSoup(archive.read(name), "html.parser")
            for node in soup(["script", "style", "noscript"]):
                node.decompose()
            text = clean_text(soup.get_text("\n", strip=True))
            if text:
                out.append(text)
    return clean_text("\n\n".join(out)), {"htmlDocuments": len(out)}



def pdf_native_metrics(path):
    import fitz
    doc = fitz.open(path)
    page_count = doc.page_count
    pages = 0
    blocks = []
    for index, page in enumerate(doc):
        text = clean_text(page.get_text("text", sort=True))
        if text:
            pages += 1
            blocks.append(f"## Page {index + 1}\n\n{text}")
    doc.close()
    return clean_text("\n\n".join(blocks)), page_count, pages


def extract_opendataloader(path, workdir):
    output = Path(workdir) / ("opendataloader-" + uuid.uuid4().hex)
    output.mkdir(parents=True, exist_ok=False)
    run_bounded(
        [
            "opendataloader-pdf",
            str(path),
            "--output-dir", str(output),
            "--format", "markdown,json",
            "--image-output", "off",
            "--threads", "1",
            "--quiet",
        ],
        timeout=OPENDATALOADER_TIMEOUT_SEC,
        label="opendataloader",
    )
    markdown_files = sorted(output.rglob("*.md"))
    if not markdown_files:
        raise RuntimeError("opendataloader_markdown_missing")
    text = clean_text(markdown_files[0].read_text("utf-8", errors="replace"))
    _, page_count, native_pages_with_text = pdf_native_metrics(path)
    pages = set()
    json_files = sorted(output.rglob("*.json"))
    if json_files:
        try:
            payload = json.loads(json_files[0].read_text("utf-8", errors="replace"))
            elements = payload if isinstance(payload, list) else payload.get("elements", payload.get("data", []))
            if isinstance(elements, list):
                for item in elements:
                    if not isinstance(item, dict):
                        continue
                    page = item.get("page number", item.get("page_number", item.get("page")))
                    if isinstance(page, int):
                        pages.add(page)
        except Exception:
            pass
    pages_with_text = len(pages) if pages else native_pages_with_text
    return text, page_count, pages_with_text

def ocr_pdf(path, workdir):
    output = Path(workdir) / ("ocr-" + uuid.uuid4().hex + ".pdf")
    command = [
        "ocrmypdf", "--force-ocr", "--rotate-pages", "--deskew",
        "--optimize", "1", "--jobs", "1", "-l", OCR_LANGUAGES,
        str(path), str(output),
    ]
    run_bounded(command, timeout=OCR_TIMEOUT_SEC, label="ocrmypdf", accepted_codes=(0, 6))
    if not output.exists():
        raise RuntimeError("ocrmypdf_output_missing")
    return output


def extract_image(path):
    stdout, _ = run_bounded(
        ["tesseract", str(path), "stdout", "-l", OCR_LANGUAGES, "--psm", "6"],
        timeout=IMAGE_OCR_TIMEOUT_SEC,
        label="tesseract",
    )
    return clean_text(stdout), {}

def libreoffice_convert(path, workdir, target, label):
    output = Path(workdir) / ("libreoffice-" + uuid.uuid4().hex)
    output.mkdir(parents=True, exist_ok=False)
    profile = Path(workdir) / ("libreoffice-profile-" + uuid.uuid4().hex)
    profile.mkdir(parents=True, exist_ok=False)
    profile_uri = profile.resolve().as_uri()
    run_bounded(
        [
            "libreoffice",
            "--headless",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to", target,
            "--outdir", str(output),
            str(path),
        ],
        timeout=OFFICE_TIMEOUT_SEC,
        label=label,
    )
    candidates = list(output.glob(f"*.{target}"))
    if not candidates:
        raise RuntimeError(f"{label}_output_missing")
    return candidates[0]


def convert_office_to_pdf(path, workdir):
    return libreoffice_convert(path, workdir, "pdf", "libreoffice_pdf_conversion")


def convert_legacy(path, workdir):
    target = LEGACY_TARGET.get(path.suffix.lower())
    if not target:
        raise RuntimeError("legacy_conversion_not_supported")
    return libreoffice_convert(path, workdir, target, "libreoffice_conversion")

def extract_pdf(path, workdir, attempts):
    try:
        text, page_count, pages = extract_opendataloader(path, workdir)
        metrics = quality(text, page_count=page_count, pages_with_text=pages)
        attempts.append({"parser": "OPENDATALOADER_LOCAL", **metrics})
        if metrics["accepted"]:
            return text, "OPENDATALOADER_LOCAL", metrics, attempts
    except Exception as exc:
        attempts.append({"parser": "OPENDATALOADER_LOCAL", "accepted": False, "error": str(exc)[:1000]})

    try:
        text, page_count, pages = pdf_native_metrics(path)
        metrics = quality(text, page_count=page_count, pages_with_text=pages)
        attempts.append({"parser": "PYMUPDF_NATIVE", **metrics})
        if metrics["accepted"]:
            return text, "PYMUPDF_NATIVE", metrics, attempts
    except Exception as exc:
        attempts.append({"parser": "PYMUPDF_NATIVE", "accepted": False, "error": str(exc)[:1000]})

    try:
        ocr_path = ocr_pdf(path, workdir)
        text, page_count, pages = extract_opendataloader(ocr_path, workdir)
        metrics = quality(text, page_count=page_count, pages_with_text=pages)
        attempts.append({"parser": "OCRMY_PDF_OPENDATALOADER", **metrics})
        if metrics["accepted"]:
            return text, "OCRMY_PDF_OPENDATALOADER", metrics, attempts
    except Exception as exc:
        attempts.append({"parser": "OCRMY_PDF_OPENDATALOADER", "accepted": False, "error": str(exc)[:1000]})

    # OCR may succeed even when the structural parser fails. The OCR-produced
    # searchable PDF is still suitable for the same native quality gate.
    if "ocr_path" in locals():
        try:
            text, page_count, pages = pdf_native_metrics(ocr_path)
            metrics = quality(text, page_count=page_count, pages_with_text=pages)
            attempts.append({"parser": "OCRMY_PDF_PYMUPDF_NATIVE", **metrics})
            if metrics["accepted"]:
                return text, "OCRMY_PDF_PYMUPDF_NATIVE", metrics, attempts
        except Exception as exc:
            attempts.append({"parser": "OCRMY_PDF_PYMUPDF_NATIVE", "accepted": False, "error": str(exc)[:1000]})

    raise RuntimeError("pdf_extraction_quality_failed:" + json.dumps(attempts, ensure_ascii=False))


def parse_document(path):
    suffix = path.suffix.lower()
    attempts = []
    with tempfile.TemporaryDirectory(prefix="gbap-extract-") as workdir:
        working = path
        if suffix in ARCHIVE_EXTENSIONS:
            validate_archive_container(working)
        if suffix in LEGACY_TARGET:
            working = convert_legacy(path, workdir)
            suffix = working.suffix.lower()
            if suffix in ARCHIVE_EXTENSIONS:
                validate_archive_container(working)
            attempts.append({"parser": "LIBREOFFICE_CONVERSION", "accepted": True})

        if suffix == ".pdf":
            return extract_pdf(working, workdir, attempts)

        if suffix == ".docx":
            text, extra = extract_docx(working)
            parser = "PYTHON_DOCX"
        elif suffix == ".xlsx":
            text, extra = extract_xlsx(working)
            parser = "OPENPYXL"
        elif suffix == ".pptx":
            text, extra = extract_pptx(working)
            parser = "PYTHON_PPTX"
        elif suffix in IMAGE_EXTENSIONS:
            text, extra = extract_image(working)
            parser = "TESSERACT_OCR"
        elif suffix in {".html", ".htm"}:
            text, extra = extract_html(working)
            parser = "BEAUTIFULSOUP"
        elif suffix == ".eml":
            text, extra = extract_eml(working)
            parser = "EMAIL_NATIVE"
        elif suffix == ".epub":
            text, extra = extract_epub(working)
            parser = "EPUB_NATIVE"
        elif suffix in TEXT_EXTENSIONS:
            text, extra = decode_text_file(working), {}
            parser = "TEXT_NATIVE"
        else:
            raise RuntimeError("unsupported_document_format")

        metrics = quality(text)
        metrics.update(extra)
        attempts.append({"parser": parser, **metrics})
        if metrics["accepted"]:
            return clean_text(text), parser, metrics, attempts

        if suffix in {".docx", ".xlsx", ".pptx"}:
            try:
                pdf_path = convert_office_to_pdf(working, workdir)
                text, pdf_parser, pdf_metrics, attempts = extract_pdf(pdf_path, workdir, attempts)
                return text, "LIBREOFFICE_PDF_" + pdf_parser, pdf_metrics, attempts
            except Exception as exc:
                attempts.append({"parser": "LIBREOFFICE_PDF_FALLBACK", "accepted": False, "error": str(exc)[:1000]})

        raise RuntimeError("document_extraction_quality_failed:" + json.dumps(attempts, ensure_ascii=False))

def process_job(job_id):
    conn = db()
    row = conn.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
    if not row or row["state"] != "QUEUED":
        conn.close()
        return
    cancel_event = threading.Event()
    with ACTIVE_JOBS_LOCK:
        ACTIVE_JOBS[job_id] = cancel_event
    claimed = conn.execute("UPDATE jobs SET state='RUNNING',stage='EXTRACTING',updated_at=? WHERE id=? AND state='QUEUED'", (now(), job_id)).rowcount
    conn.commit()
    conn.close()
    if not claimed:
        with ACTIVE_JOBS_LOCK:
            ACTIVE_JOBS.pop(job_id, None)
        return
    CURRENT_JOB.cancel_event = cancel_event

    job_dir = JOBS_ROOT / job_id
    input_path = Path(row["input_path"])
    result_path = job_dir / "normalized.md"
    try:
        text, parser, metrics, attempts = parse_document_isolated(input_path)
        if len(text) > MAX_NORMALIZED_CHARS:
            raise RuntimeError(f"normalized_content_too_large:{len(text)}>{MAX_NORMALIZED_CHARS}")
        text = f"<!-- GBAP_NORMALIZED parser={parser} -->\n\n{text}\n"
        result_path.write_text(text, encoding="utf-8")
        conn = db()
        current = conn.execute("SELECT state FROM jobs WHERE id=?", (job_id,)).fetchone()
        if current and current["state"] == "CANCELLED":
            conn.close()
            shutil.rmtree(job_dir, ignore_errors=True)
            return
        metrics["attempts"] = attempts
        conn.execute(
            "UPDATE jobs SET state='READY',stage='NORMALIZED',result_path=?,parser=?,metrics_json=?,error=NULL,updated_at=? WHERE id=?",
            (str(result_path), parser, json.dumps(metrics, ensure_ascii=False), now(), job_id),
        )
        conn.commit()
        conn.close()
    except Exception as exc:
        conn = db()
        current = conn.execute("SELECT state FROM jobs WHERE id=?", (job_id,)).fetchone()
        if current and current["state"] != "CANCELLED" and not STOP.is_set():
            conn.execute(
                "UPDATE jobs SET state='FAILED',stage='FAILED',error=?,updated_at=? WHERE id=?",
                (str(exc)[:8000], now(), job_id),
            )
            conn.commit()
        conn.close()
    finally:
        with ACTIVE_JOBS_LOCK:
            ACTIVE_JOBS.pop(job_id, None)
        CURRENT_JOB.cancel_event = None
        if cancel_event.is_set():
            shutil.rmtree(job_dir, ignore_errors=True)


def worker():
    last_maintenance = 0.0
    while not STOP.is_set():
        current = time.monotonic()
        if current - last_maintenance >= MAINTENANCE_INTERVAL_SEC:
            try:
                purged = purge_expired_jobs()
                if purged:
                    print(json.dumps({"event": "knowledge-extractor.maintenance", "purgedJobs": purged}), flush=True)
            except Exception as exc:
                print(json.dumps({"event": "knowledge-extractor.maintenance_failed", "error": str(exc)[:1000]}), flush=True)
            last_maintenance = current
        conn = db()
        row = conn.execute("SELECT id FROM jobs WHERE state='QUEUED' ORDER BY created_at LIMIT 1").fetchone()
        conn.close()
        if row:
            process_job(row["id"])
            continue
        WAKE.wait(2)
        WAKE.clear()


def public_job(row):
    return {
        "jobId": row["id"],
        "state": row["state"],
        "stage": row["stage"],
        "sourceFilename": row["source_filename"],
        "parser": row["parser"],
        "metrics": json.loads(row["metrics_json"]) if row["metrics_json"] else None,
        "error": row["error"],
        "handoffExternalId": row["handoff_external_id"],
        "createdAt": row["created_at"],
        "updatedAt": row["updated_at"],
    }


class Handler(BaseHTTPRequestHandler):
    server_version = "GBAPKnowledgeExtractor/1.0"

    def log_message(self, fmt, *args):
        print(json.dumps({"event": "knowledge-extractor.http", "message": fmt % args}, ensure_ascii=False), flush=True)

    def send_json(self, status, payload):
        raw = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("content-type", "application/json; charset=utf-8")
        self.send_header("cache-control", "no-store")
        self.send_header("content-length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def read_json(self, max_bytes):
        length = int(self.headers.get("content-length", "0") or "0")
        if length <= 0 or length > max_bytes:
            raise ValueError("invalid_request_size")
        return json.loads(self.rfile.read(length))

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/health":
            return self.send_json(200, {"status": "ok", "service": "knowledge-extractor", "version": "1.0"})
        match = re.fullmatch(r"/jobs/([0-9a-f\-]{36})(/result)?", path)
        if not match:
            return self.send_json(404, {"error": "not_found"})
        conn = db()
        row = conn.execute("SELECT * FROM jobs WHERE id=?", (match.group(1),)).fetchone()
        conn.close()
        if not row:
            return self.send_json(404, {"error": "job_not_found"})
        if match.group(2):
            if row["state"] != "READY" or not row["result_path"] or not Path(row["result_path"]).exists():
                return self.send_json(409, {"error": "result_not_ready"})
            raw = Path(row["result_path"]).read_bytes()
            return self.send_json(200, {
                "jobId": row["id"],
                "sourceFilename": row["source_filename"],
                "normalizedFilename": f"gbap-{row['id']}.md",
                "mediaType": "text/markdown",
                "contentBase64": base64.b64encode(raw).decode("ascii"),
                "parser": row["parser"],
                "metrics": json.loads(row["metrics_json"]) if row["metrics_json"] else None,
            })
        return self.send_json(200, public_job(row))

    def do_POST(self):
        path = urlparse(self.path).path
        if path != "/jobs":
            return self.send_json(404, {"error": "not_found"})
        try:
            payload = self.read_json(int(MAX_BYTES * 1.45) + 256 * 1024)
            filename = safe_name(str(payload.get("sourceFilename") or "document.bin"))
            suffix = Path(filename).suffix.lower()
            if suffix not in SUPPORTED:
                return self.send_json(415, {"error": "unsupported_document_format", "extension": suffix})
            encoded = str(payload.get("contentBase64") or "")
            raw = base64.b64decode(encoded, validate=True)
            if not raw or len(raw) > MAX_BYTES:
                return self.send_json(413 if len(raw) > MAX_BYTES else 400, {"error": "invalid_document_size"})
            job_id = str(uuid.uuid4())
            job_dir = JOBS_ROOT / job_id
            job_dir.mkdir(parents=True, exist_ok=False)
            input_path = job_dir / ("source" + suffix)
            input_path.write_bytes(raw)
            at = now()
            conn = db()
            conn.execute(
                "INSERT INTO jobs(id,state,stage,source_filename,media_type,input_path,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",
                (job_id, "QUEUED", "QUEUED", filename, str(payload.get("mediaType") or "application/octet-stream"), str(input_path), at, at),
            )
            conn.commit()
            row = conn.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
            conn.close()
            WAKE.set()
            return self.send_json(202, public_job(row))
        except Exception as exc:
            return self.send_json(400, {"error": str(exc)[:1000]})

    def do_PUT(self):
        path = urlparse(self.path).path
        match = re.fullmatch(r"/jobs/([0-9a-f\-]{36})/handoff", path)
        if not match:
            return self.send_json(404, {"error": "not_found"})
        try:
            payload = self.read_json(64 * 1024)
            external_id = str(payload.get("externalId") or "").strip()
            if not external_id or len(external_id) > 255:
                return self.send_json(400, {"error": "invalid_external_id"})
            conn = db()
            row = conn.execute("SELECT * FROM jobs WHERE id=?", (match.group(1),)).fetchone()
            if not row:
                conn.close()
                return self.send_json(404, {"error": "job_not_found"})
            conn.execute("UPDATE jobs SET handoff_external_id=?,updated_at=? WHERE id=?", (external_id, now(), row["id"]))
            conn.commit()
            updated = conn.execute("SELECT * FROM jobs WHERE id=?", (row["id"],)).fetchone()
            conn.close()
            return self.send_json(200, public_job(updated))
        except Exception as exc:
            return self.send_json(400, {"error": str(exc)[:1000]})

    def do_DELETE(self):
        path = urlparse(self.path).path
        match = re.fullmatch(r"/jobs/([0-9a-f\-]{36})", path)
        if not match:
            return self.send_json(404, {"error": "not_found"})
        conn = db()
        row = conn.execute("SELECT * FROM jobs WHERE id=?", (match.group(1),)).fetchone()
        if not row:
            conn.close()
            return self.send_json(200, {"removed": True})
        if row["state"] == "RUNNING":
            conn.execute("UPDATE jobs SET state='CANCELLED',stage='CANCELLED',updated_at=? WHERE id=?", (now(), row["id"]))
        else:
            conn.execute("DELETE FROM jobs WHERE id=?", (row["id"],))
        conn.commit()
        conn.close()
        if row["state"] == "RUNNING":
            with ACTIVE_JOBS_LOCK:
                cancel_event = ACTIVE_JOBS.get(row["id"])
                if cancel_event:
                    cancel_event.set()
        if row["state"] != "RUNNING":
            shutil.rmtree(JOBS_ROOT / row["id"], ignore_errors=True)
        return self.send_json(200, {"removed": True})


if __name__ == "__main__":
    if len(sys.argv) == 4 and sys.argv[1] == "--parse-document":
        sys.exit(parser_child_main(sys.argv[2], sys.argv[3]))
    ensure_storage()
    reset_interrupted_jobs()
    thread = threading.Thread(target=worker, name="knowledge-extractor-worker", daemon=True)
    thread.start()
    server = ThreadingHTTPServer(("0.0.0.0", PORT), Handler)
    print(json.dumps({"event": "knowledge-extractor.started", "port": PORT, "dataRoot": str(DATA_ROOT)}), flush=True)
    def shutdown_service(_signal, _frame):
        STOP.set()
        WAKE.set()
        threading.Thread(target=server.shutdown, daemon=True).start()
    signal.signal(signal.SIGTERM, shutdown_service)
    try:
        server.serve_forever()
    finally:
        STOP.set()
        WAKE.set()
        thread.join(timeout=5)
        server.server_close()
